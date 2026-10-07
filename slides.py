import json
import os
import re
import time

import requests
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


# ============================================================
# SETTINGS
# ============================================================

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()

# Render'da GEMINI_MODEL bo'lmasa, shu model ishlaydi
GEMINI_MODEL = os.getenv(
    "GEMINI_MODEL",
    "gemini-2.5-flash"
).strip()


# ============================================================
# COLORS
# ============================================================

DARK = RGBColor(31, 42, 68)
ORANGE = RGBColor(242, 122, 26)
LIGHT = RGBColor(247, 245, 240)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(208, 213, 224)


# ============================================================
# AI PROMPT
# ============================================================

PROMPT = """
Sen professional PowerPoint prezentatsiya tayyorlovchi AI yordamchisan.

Foydalanuvchi so'rovi:
__REQUEST__

Vazifa:
Ushbu mavzu asosida professional prezentatsiya yarat.

Qoidalar:

- Foydalanuvchi qaysi tilda yozgan bo'lsa, shu tilda yoz.
- Agar foydalanuvchi slayd sonini ko'rsatgan bo'lsa, aynan shuncha slayd yarat.
- Agar son ko'rsatilmagan bo'lsa, 7 ta slayd yarat.
- Minimal 3 ta, maksimal 12 ta slayd.
- Birinchi slayd titul slayd bo'lsin.
- Birinchi slaydda bullets bo'lmasin.
- Birinchi slaydda subtitle bo'lsin.
- Oxirgi slayd xulosa bo'lsin.
- Oddiy slaydlarda 3-5 ta qisqa bullet bo'lsin.
- Har bir bullet 15 so'zdan oshmasin.
- Ma'lumotlar tushunarli va mantiqiy bo'lsin.
- Bir xil ma'lumotni takrorlama.

FAQAT JSON qaytar.
Markdown ishlatma.
```json ishlatma.
Hech qanday qo'shimcha izoh yozma.

JSON formati:

{
  "title": "Prezentatsiya nomi",
  "slides": [
    {
      "title": "Kirish",
      "subtitle": "Qisqa izoh",
      "bullets": []
    },
    {
      "title": "Asosiy mavzu",
      "subtitle": "",
      "bullets": [
        "Birinchi fikr",
        "Ikkinchi fikr",
        "Uchinchi fikr"
      ]
    }
  ]
}
"""


# ============================================================
# GEMINI
# ============================================================

def ask_gemini(request_text: str) -> dict:

    if not GEMINI_API_KEY:
        raise RuntimeError(
            "GEMINI_API_KEY topilmadi. "
            "Render Environment bo'limini tekshiring."
        )

    # .format() ISHLATILMAYDI!
    # Shuning uchun JSON {} sababli KeyError bo'lmaydi.
    prompt = PROMPT.replace(
        "__REQUEST__",
        request_text
    )

    url = (
        "https://generativelanguage.googleapis.com/v1beta/models/"
        + GEMINI_MODEL
        + ":generateContent"
    )

    payload = {
        "contents": [
            {
                "parts": [
                    {
                        "text": prompt
                    }
                ]
            }
        ],
        "generationConfig": {
            "temperature": 0.7,
            "responseMimeType": "application/json"
        }
    }

    last_error = None

    # Gemini vaqtincha band bo'lsa 4 marta urinadi
    for attempt in range(4):

        try:

            response = requests.post(
                url,
                params={
                    "key": GEMINI_API_KEY
                },
                json=payload,
                timeout=180
            )

            # Vaqtinchalik server xatolari
            if response.status_code in (
                429,
                500,
                502,
                503,
                504
            ):

                last_error = (
                    f"Gemini HTTP {response.status_code}"
                )

                if attempt < 3:
                    time.sleep(3 * (attempt + 1))
                    continue

                raise RuntimeError(
                    "Gemini serveri vaqtincha band. "
                    "Bir necha daqiqadan keyin qayta urinib ko'ring."
                )

            response.raise_for_status()

            result = response.json()

            candidates = result.get(
                "candidates",
                []
            )

            if not candidates:
                raise RuntimeError(
                    "Gemini javob qaytarmadi."
                )

            content = candidates[0].get(
                "content",
                {}
            )

            parts = content.get(
                "parts",
                []
            )

            if not parts:
                raise RuntimeError(
                    "Gemini javobida matn topilmadi."
                )

            text = parts[0].get(
                "text",
                ""
            ).strip()

            if not text:
                raise RuntimeError(
                    "Gemini bo'sh javob qaytardi."
                )

            # ==================================================
            # CLEAN JSON
            # ==================================================

            text = re.sub(
                r"^```json\s*",
                "",
                text,
                flags=re.IGNORECASE
            )

            text = re.sub(
                r"^```\s*",
                "",
                text
            )

            text = re.sub(
                r"\s*```$",
                "",
                text
            )

            text = text.strip()

            # JSON parse
            try:

                data = json.loads(text)

            except json.JSONDecodeError:

                # Agar Gemini ortiqcha matn qo'shgan bo'lsa
                start = text.find("{")
                end = text.rfind("}")

                if start == -1 or end == -1:
                    raise RuntimeError(
                        "Gemini valid JSON qaytarmadi."
                    )

                data = json.loads(
                    text[start:end + 1]
                )

            # ==================================================
            # VALIDATION
            # ==================================================

            if not isinstance(data, dict):
                raise RuntimeError(
                    "Gemini noto'g'ri format qaytardi."
                )

            slides = data.get(
                "slides",
                []
            )

            if not isinstance(slides, list):
                raise RuntimeError(
                    "Slaydlar ro'yxati topilmadi."
                )

            if len(slides) == 0:
                raise RuntimeError(
                    "Gemini hech qanday slayd yaratmadi."
                )

            # Maksimum 12 ta
            slides = slides[:12]

            # Har bir slaydni standartlashtiramiz
            cleaned_slides = []

            for slide in slides:

                if not isinstance(slide, dict):
                    continue

                title = str(
                    slide.get(
                        "title",
                        "Slayd"
                    )
                )

                subtitle = str(
                    slide.get(
                        "subtitle",
                        ""
                    )
                )

                bullets = slide.get(
                    "bullets",
                    []
                )

                if not isinstance(
                    bullets,
                    list
                ):
                    bullets = []

                bullets = [
                    str(x).strip()
                    for x in bullets
                    if str(x).strip()
                ]

                cleaned_slides.append({
                    "title": title,
                    "subtitle": subtitle,
                    "bullets": bullets
                })

            if not cleaned_slides:
                raise RuntimeError(
                    "Slaydlar qayta ishlanmadi."
                )

            return {
                "title": str(
                    data.get(
                        "title",
                        "Taqdimot"
                    )
                ),
                "slides": cleaned_slides
            }

        except requests.RequestException as error:

            last_error = str(error)

            if attempt < 3:
                time.sleep(3 * (attempt + 1))
                continue

            raise RuntimeError(
                f"Gemini bilan bog'lanishda xatolik: {last_error}"
            )

    raise RuntimeError(
        last_error or "Noma'lum Gemini xatosi."
    )


# ============================================================
# POWERPOINT FUNCTIONS
# ============================================================

def fill_background(slide, color):

    fill = slide.background.fill

    fill.solid()

    fill.fore_color.rgb = color


def add_text(
    slide,
    x,
    y,
    width,
    height,
    text,
    size,
    color,
    bold=False
):

    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )

    frame = box.text_frame

    frame.word_wrap = True

    paragraph = frame.paragraphs[0]

    paragraph.text = str(text)

    paragraph.font.size = Pt(size)

    paragraph.font.bold = bold

    paragraph.font.color.rgb = color

    return box


# ============================================================
# BUILD PPTX
# ============================================================

def build_pptx(data: dict, path: str) -> str:

    presentation = Presentation()

    # 16:9
    presentation.slide_width = Inches(13.333)

    presentation.slide_height = Inches(7.5)

    blank = presentation.slide_layouts[6]

    slides = data.get(
        "slides",
        []
    )

    presentation_title = data.get(
        "title",
        "Taqdimot"
    )

    for index, slide_data in enumerate(slides):

        slide = presentation.slides.add_slide(
            blank
        )

        title = slide_data.get(
            "title",
            "Slayd"
        )

        subtitle = slide_data.get(
            "subtitle",
            ""
        )

        bullets = slide_data.get(
            "bullets",
            []
        )

        # ====================================================
        # TITLE SLIDE
        # ====================================================

        if index == 0:

            fill_background(
                slide,
                DARK
            )

            # Orange line
            line = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.8),
                Inches(3.0),
                Inches(1.7),
                Inches(0.12)
            )

            line.fill.solid()

            line.fill.fore_color.rgb = ORANGE

            line.line.fill.background()

            # Main title
            add_text(
                slide,
                0.8,
                1.25,
                11.5,
                1.6,
                presentation_title or title,
                44,
                WHITE,
                True
            )

            # Subtitle
            add_text(
                slide,
                0.8,
                3.35,
                11.5,
                1.3,
                subtitle or title,
                23,
                GRAY
            )

            continue

        # ====================================================
        # NORMAL SLIDE
        # ====================================================

        fill_background(
            slide,
            LIGHT
        )

        # Orange side bar
        side = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0,
            0,
            Inches(0.32),
            presentation.slide_height
        )

        side.fill.solid()

        side.fill.fore_color.rgb = ORANGE

        side.line.fill.background()

        # Title
        add_text(
            slide,
            0.85,
            0.45,
            11.3,
            1.0,
            title,
            34,
            DARK,
            True
        )

        # Bullet container
        box = slide.shapes.add_textbox(
            Inches(0.9),
            Inches(1.75),
            Inches(11.4),
            Inches(4.9)
        )

        frame = box.text_frame

        frame.word_wrap = True

        for i, bullet in enumerate(bullets):

            paragraph = (
                frame.paragraphs[0]
                if i == 0
                else frame.add_paragraph()
            )

            paragraph.text = (
                "•  " + str(bullet)
            )

            paragraph.font.size = Pt(23)

            paragraph.font.color.rgb = DARK

            paragraph.space_after = Pt(15)

        # Slide number
        add_text(
            slide,
            12.0,
            6.85,
            0.6,
            0.3,
            str(index + 1),
            14,
            ORANGE,
            True
        )

    # ========================================================
    # SAVE
    # ========================================================

    presentation.save(path)

    return path
