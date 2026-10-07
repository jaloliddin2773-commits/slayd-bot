import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DARK_BLUE = RGBColor(15, 23, 42)     # #0F172A
ACCENT_BLUE = RGBColor(37, 99, 235)  # #2563EB
TEXT_DARK = RGBColor(51, 65, 85)      # #334155
WHITE = RGBColor(255, 255, 255)

def create_pptx(slide_data: dict) -> io.BytesIO:
    prs = Presentation()
    # Professional 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]

    # Cover Slide
    cover_slide = prs.slides.add_slide(blank_layout)
    bg_shape = cover_slide.shapes.add_shape(
        1, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = DARK_BLUE
    bg_shape.line.fill.background()

    tx_box = cover_slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(2.5))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = slide_data.get("title", "Prezentatsiya")
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    # Content Slides
    for slide_info in slide_data.get("slides", []):
        slide = prs.slides.add_slide(blank_layout)
        
        # Header Line
        header_line = slide.shapes.add_shape(
            1, Inches(0.8), Inches(0.8), Inches(0.15), Inches(0.8)
        )
        header_line.fill.solid()
        header_line.fill.fore_color.rgb = ACCENT_BLUE
        header_line.line.fill.background()

        # Title
        title_box = slide.shapes.add_textbox(Inches(1.2), Inches(0.6), Inches(11), Inches(1))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = slide_info.get("title", "")
        p_title.font.bold = True
        p_title.font.size = Pt(28)
        p_title.font.color.rgb = DARK_BLUE

        # Bullet Points
        content_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11), Inches(4.5))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True

        for i, point in enumerate(slide_info.get("bullet_points", [])):
            p_point = tf_content.add_paragraph() if i > 0 else tf_content.paragraphs[0]
            p_point.text = f"• {point}"
            p_point.font.size = Pt(18)
            p_point.font.color.rgb = TEXT_DARK
            p_point.space_after = Pt(14)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output
